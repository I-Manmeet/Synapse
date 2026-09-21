from dotenv import load_dotenv
from azure.identity import DefaultAzureCredential
from azure.storage.blob import BlobServiceClient
from pypdf import PdfReader
import os
import pandas as pd
from io import BytesIO
from docx import Document

load_dotenv()

STORAGE_ACCOUNT_URL = os.getenv(
    "AZURE_STORAGE_ACCOUNT_URL"
)

CONTAINER_NAME = "synapse-documents"

if not STORAGE_ACCOUNT_URL:
    raise ValueError(
        "AZURE_STORAGE_ACCOUNT_URL is missing from .env"
    )

credential = DefaultAzureCredential()

blob_service_client = BlobServiceClient(
    account_url=STORAGE_ACCOUNT_URL,
    credential=credential
)


def download_blob(filename: str) -> bytes:

    blob_client = blob_service_client.get_blob_client(
        container=CONTAINER_NAME,
        blob=filename
    )

    return blob_client.download_blob().readall()


def extract_text_from_txt(filename: str) -> str:

    data = download_blob(filename)

    return data.decode("utf-8")


def extract_data_from_csv(filename: str):

    data = download_blob(filename)

    df = pd.read_csv(BytesIO(data))

    return df


def extract_data_from_excel(filename: str):

    data = download_blob(filename)

    df = pd.read_excel(BytesIO(data))

    return df
def extract_text_from_pdf(filename: str) -> str:

    data = download_blob(filename)

    reader = PdfReader(BytesIO(data))

    text = ""

    for i, page in enumerate(reader.pages):

        page_text = page.extract_text()

        if page_text:
            text += page_text + "\n"

    if text.strip():
        return text

    return ""

    data = download_blob(filename)

    print("PDF downloaded successfully")
    print("PDF size:", len(data), "bytes")

    reader = PdfReader(BytesIO(data))

    print("Number of pages:", len(reader.pages))

    text = ""

    for i, page in enumerate(reader.pages):

        print(f"Processing page {i + 1}...")

        page_text = page.extract_text()

        print("Extracted text:", repr(page_text))

        if page_text:
            text += page_text + "\n"

    return text

    data = download_blob(filename)

    reader = PdfReader(BytesIO(data))

    text = ""

    for page in reader.pages:
        page_text = page.extract_text()

        if page_text:
            text += page_text + "\n"

    return text

def extract_text_from_docx(filename: str) -> str:

    data = download_blob(filename)

    document = Document(BytesIO(data))

    text = ""

    # Extract normal paragraphs
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text += paragraph.text + "\n"

    # Extract tables
    for table in document.tables:
        for row in table.rows:
            row_data = []

            for cell in row.cells:
                row_data.append(cell.text.strip())

            text += " | ".join(row_data) + "\n"

    return text